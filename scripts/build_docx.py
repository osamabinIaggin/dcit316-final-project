#!/usr/bin/env python3
"""Render the paper to .docx for Sakai submission.

    python3 scripts/build_docx.py [source.md]

Handles the subset of Markdown the paper uses: ATX headings, pipe tables,
blockquotes, bullet lists and inline bold/italic/code.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DEFAULT = ROOT / "DCIT316_Audit_22128981_GideonElormGlago.md"

INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`|\[.+?\]\(.+?\))")


def add_runs(para, text: str) -> None:
    """Split text on inline markers and style each fragment."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            para.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            para.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = para.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xA0, 0x30, 0x30)
        elif part.startswith("["):
            label, _, url = part[1:].partition("](")
            run = para.add_run(label)
            run.font.color.rgb = RGBColor(0x18, 0x50, 0xA0)
            run.underline = True
            if url.rstrip(")") not in label:
                para.add_run(f" ({url.rstrip(')')})").font.size = Pt(8)
        else:
            para.add_run(part)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build(src: Path, out: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    lines = src.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip() or line.strip() in {"---", "***"}:
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            header = split_row(line)
            i += 2
            body = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                body.append(split_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, header):
                cell.paragraphs[0].runs.clear() if cell.paragraphs[0].runs else None
                add_runs(cell.paragraphs[0], text or " ")
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in body:
                cells = table.add_row().cells
                for cell, text in zip(cells, row + [""] * (len(header) - len(row))):
                    add_runs(cell.paragraphs[0], text or " ")
                    for r in cell.paragraphs[0].runs:
                        r.font.size = Pt(9.5)
            doc.add_paragraph()
            continue

        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("# ").strip()
            if level == 1 and not doc.paragraphs:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text)
                run.bold = True
                run.font.size = Pt(19)
            else:
                add_runs(doc.add_heading(level=min(level, 4)), text)
            i += 1
            continue

        if line.startswith(">"):
            block = []
            while i < len(lines) and lines[i].startswith(">"):
                block.append(lines[i].lstrip("> ").rstrip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(26)
            add_runs(p, " ".join(b for b in block if b))
            for r in p.runs:
                r.italic = True
            continue

        if re.match(r"^\s*[-*]\s+", line):
            add_runs(doc.add_paragraph(style="List Bullet"),
                     re.sub(r"^\s*[-*]\s+", "", line))
            i += 1
            continue

        if re.match(r"^\s*\d+\.\s+", line):
            add_runs(doc.add_paragraph(style="List Number"),
                     re.sub(r"^\s*\d+\.\s+", "", line))
            i += 1
            continue

        para = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^(#|\||>|\s*[-*]\s|\s*\d+\.\s|---)", lines[i]):
            para.append(lines[i].rstrip())
            i += 1
        add_runs(doc.add_paragraph(), " ".join(para))

    doc.save(out)


def main() -> int:
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT
    out = src.with_suffix(".docx")
    build(src, out)
    print(f"{out.name}  ({out.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
